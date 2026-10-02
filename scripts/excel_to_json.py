"""
Coffee Atlas V1.6
Excel -> JSON converter

Human-maintained source:
    Coffee_Atlas_Database_V1.6.xlsx

Website output:
    data/atlas.json
"""

from pathlib import Path
import json
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EXCEL_PATH = ROOT / "Coffee_Atlas_Database_V1.6.xlsx"
OUTPUT_PATH = ROOT / "data" / "atlas.json"


def clean(value):
    """Return a trimmed value; Excel blank cells become None."""
    if value is None:
        return None
    if isinstance(value, str):
        value = value.strip()
        return value if value else None
    return value


def split_list(value):
    """Convert 'Floral, Citrus, Sweet' into a JSON list."""
    value = clean(value)
    if not value:
        return []
    return [item.strip() for item in str(value).split(",") if item.strip()]


def sheet_rows(ws):
    """Read an Excel sheet as dictionaries using row 1 as headers."""
    headers = [clean(cell.value) for cell in ws[1]]
    rows = []

    for values in ws.iter_rows(min_row=2, values_only=True):
        row = {
            headers[index]: clean(value)
            for index, value in enumerate(values)
            if index < len(headers) and headers[index]
        }

        # Completely blank rows are ignored.
        if any(value is not None for value in row.values()):
            rows.append(row)

    return rows


def active(row):
    """Draft rows stay in Excel but do not go live on the website."""
    return row.get("Status") in (None, "Active", "Legacy Demo")


def score(value):
    """Keep sensory scores optional instead of inventing 0/5."""
    if value in (None, ""):
        return None
    try:
        return int(value)
    except (TypeError, ValueError):
        return value


def main():
    wb = load_workbook(EXCEL_PATH, data_only=True)

    country_rows = sheet_rows(wb["Countries"])
    region_rows = sheet_rows(wb["Regions"])
    farm_rows = sheet_rows(wb["Farms"])
    bean_rows = sheet_rows(wb["Beans"])

    countries = []
    for row in country_rows:
        if not active(row):
            continue
        countries.append({
            "id": row.get("Country ID"),
            "name": row.get("Country EN"),
            "zh": row.get("Country ZH"),
            "profile": split_list(row.get("Typical Profile")),
        })

    regions = []
    for row in region_rows:
        if not active(row):
            continue
        regions.append({
            "id": row.get("Region ID"),
            "country": row.get("Country ID"),
            "name": row.get("Region EN"),
            "zh": row.get("Region ZH"),
            "altitude": row.get("Altitude"),
            "profile": split_list(row.get("Typical Profile")),
        })

    stations = []
    for row in farm_rows:
        if not active(row):
            continue
        stations.append({
            "id": row.get("Farm ID"),
            "region": row.get("Region ID"),
            "name": row.get("Farm / Station EN"),
            "zh": row.get("Farm / Station ZH"),
            "type": row.get("Type"),
            "altitude": row.get("Altitude"),
            "profile": split_list(row.get("Typical Profile")),
            "intro": row.get("Intro") or "",
        })

    beans = []
    for row in bean_rows:
        if not active(row):
            continue

        flavors = [
            row.get("Flavor 1"),
            row.get("Flavor 2"),
            row.get("Flavor 3"),
            row.get("Flavor 4"),
        ]

        bean = {
            "id": row.get("Bean ID"),
            "station": row.get("Farm ID"),
            "name": row.get("Display Name"),
            "process": row.get("Process"),
            "variety": row.get("Variety"),
            "altitude": row.get("Altitude"),
            "flavors": [f for f in flavors if f],
            "sweetness": score(row.get("Sweetness (1-5)")),
            "acidity": score(row.get("Acidity (1-5)")),
            "body": score(row.get("Body (1-5)")),
        }

        # Extra metadata can exist without breaking the current UI.
        if row.get("Lot / Year"):
            bean["lot"] = row["Lot / Year"]
        if row.get("Notes"):
            bean["note"] = row["Notes"]
        if row.get("Source URL"):
            bean["source"] = row["Source URL"]

        beans.append(bean)

    # Learn remains static for now; it can become another Excel sheet later.
    brews = [
        {
            "name": "Pour Over 手沖",
            "temp": "90–94°C",
            "ratio": "15g : 240g",
            "time": "2:30–3:00",
            "roast": "淺焙／中焙",
        },
        {
            "name": "French Press 法壓",
            "temp": "88–94°C",
            "ratio": "15g : 240g",
            "time": "4:00",
            "roast": "中焙／深焙",
        },
        {
            "name": "AeroPress",
            "temp": "85–92°C",
            "ratio": "15g : 200–230g",
            "time": "1:30–2:30",
            "roast": "全烘焙度",
        },
    ]

    data = {
        "countries": countries,
        "regions": regions,
        "stations": stations,
        "beans": beans,
        "brews": brews,
    }

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text(
        json.dumps(data, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )

    print(
        "Generated atlas.json:",
        f"{len(countries)} countries,",
        f"{len(regions)} regions,",
        f"{len(stations)} farms/stations,",
        f"{len(beans)} beans.",
    )


if __name__ == "__main__":
    main()
